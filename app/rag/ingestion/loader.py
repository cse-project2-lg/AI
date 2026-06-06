from pathlib import Path
from typing import Optional
from docx import Document as DocxDocument

from app.rag.schemas.document import Document


SUPPORTED_EXTENSIONS = {".txt", ".docx"}


# 파일 경로를 받아서 Document 객체로 반환하는 함수
def load_document(file_path: str, doc_type: Optional[str] = None) -> Document:
    """
    Args:
        file_path: 원본 문서 경로
        doc_type: 문서 유형. 예: srs, policy, rule, state
    """

    path = Path(file_path)

    _validate_file(path)

    if path.suffix.lower() == ".txt":
        text = _load_txt(path)
    elif path.suffix.lower() == ".docx":
        text = _load_docx(path)
    else:
        raise ValueError(f"지원하지 않는 파일 형식입니다: {path.suffix}")

    return Document(
        doc_id=_build_doc_id(path),
        source=str(path),
        text=text,
        title=path.stem,
        doc_type=doc_type or _infer_doc_type(path),
        extra_metadata={
            "file_name": path.name,
            "extension": path.suffix,
            "parent_dir": path.parent.name,
        },
    )


# 파일 존재 여부와 확장자를 검증
def _validate_file(path: Path) -> None:

    if not path.exists():
        raise FileNotFoundError(f"파일을 찾을 수 없습니다: {path}")

    if not path.is_file():
        raise ValueError(f"파일이 아닙니다: {path}")

    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"지원하지 않는 파일 형식입니다: {path.suffix}. "
            f"지원 형식: {', '.join(SUPPORTED_EXTENSIONS)}"
        )


def _load_txt(path: Path) -> str:

    with path.open("r", encoding="utf-8") as file:
        return file.read()


def _load_docx(path: Path) -> str:

    docx = DocxDocument(path)

    blocks = []

    for paragraph in docx.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(text)

    for table in docx.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))

    return "\n".join(blocks)

def _build_doc_id(path: Path) -> str:

    return path.stem

# 파일 경로를 기반으로 문서 유형을 추론한다.

def _infer_doc_type(path: Path) -> str:

    parent_dir = path.parent.name

    if parent_dir in {"srs", "policies", "states", "rules", "meeting_notes", "knowledge_base"}:
        return parent_dir

    return "unknown"